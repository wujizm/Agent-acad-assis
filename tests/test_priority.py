import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock

from openpyxl import load_workbook

from research_agent.collect import Collector
from research_agent.config import Config
from research_agent.export import export_topic
from research_agent.journal_metrics import journal_metric
from research_agent.screening import rank_sources, relevance, screen
from research_agent.store import Store


class PriorityTests(unittest.TestCase):
    def test_if_match_requires_exact_journal_or_issn_and_retains_year(self):
        record = {"kind": "article", "venue": "The Journal of Chemical Physics"}
        self.assertEqual(journal_metric(record)["year"], 2025)
        self.assertEqual(journal_metric(record)["value"], 3.7)
        self.assertEqual(journal_metric({**record, "venue": "Journal of Chemical Physics Education"})["status"], "未核实")
        self.assertEqual(journal_metric({**record, "kind": "book"})["status"], "不适用")
        self.assertEqual(journal_metric({**record, "record_type": "proceedings-article"})["status"], "不适用")
        self.assertEqual(journal_metric({**record, "venue": "Different title", "issn": ["0021-9606"]})["value"], 3.7)

    def test_alias_search_preserves_specificity_and_if_does_not_rescue_irrelevance(self):
        query = ["machine learning interatomic potentials", "neural network potentials"]
        self.assertGreater(relevance({"title": "Neural network potentials for metals"}, query), .65)
        record = {"title": "Optical galaxy surveys", "kind": "article", "venue": "Reviews of Modern Physics", "year": "2025"}
        audit, _ = screen([record], query, {"level": "high"}, 5, year=2026)
        self.assertFalse(audit[0]["selected"])

    def test_if_channel_ranking_and_missing_values(self):
        base = {"title": "AI for science", "kind": "article", "year": "2020", "evidence_level": "abstract"}
        rows = [{**base, "id": "high", "venue": "Reviews of Modern Physics"},
                {**base, "id": "low", "venue": "Physical Review B"},
                {**base, "id": "missing", "venue": "Unknown journal"},
                {**base, "id": "excluded", "venue": "Reviews of Modern Physics", "excluded": True}]
        ranked = rank_sources(rows, "AI for science", 2026)
        by_id = {r["id"]: r for r in ranked}
        self.assertGreater(by_id["high"]["importance"]["score"], by_id["low"]["importance"]["score"])
        self.assertIsNone(by_id["missing"]["journal_metrics"]["value"])
        self.assertNotIn("期刊IF", by_id["missing"]["importance"]["reason"])
        self.assertIsNone(by_id["excluded"]["importance"]["rank"])
        audit, _ = screen([rows[0]], "AI for science", {"level": "high"}, 5, year=2026)
        self.assertTrue(audit[0]["selected"])
        self.assertIn("IF", audit[0]["reason"])

    def test_multi_query_budget_and_auditable_plan(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(folder)
            topic = store.create_topic("机器学习势函数")["id"]
            llm = Mock(ready=True)
            llm.complete.return_value = {"queries": ["machine learning interatomic potentials", "neural network potentials"]}
            net = Mock()
            net.json.side_effect = lambda url, params: {"message": {"items": []}} if "crossref" in url else {"docs": []}
            result = Collector(Config(Path(folder)), net, store, llm).run(topic, 10, lambda _: None)
            calls = [call.args[1] for call in net.json.call_args_list if "crossref" in call.args[0]]
            self.assertEqual(sum(c["rows"] for c in calls), 200)
            self.assertTrue(all("query.title" in c and "query.bibliographic" not in c for c in calls))
            self.assertEqual(len(result["queries"]), 2)
            self.assertEqual(json.loads((store.topic_dir(topic) / result["audit_file"]).read_text(encoding="utf-8"))["queries"], result["queries"])

    def test_semantic_relevance_rescues_synonyms_and_rejects_word_overlap(self):
        queries = ["electrical transport metals"]
        llm = Mock(ready=True)
        llm.complete.return_value = {"items": [
            {"index": 0, "score": 1, "reason": "量子线性响应输运理论"},
            {"index": 1, "score": 0, "reason": "金属添加剂的聚合物应用，非金属导电理论"}]}
        records = [{"title": "Kubo formula and linear response", "kind": "article", "year": "2020"},
                   {"title": "Electrical transport of metal doped polymers", "kind": "article", "year": "2020"}]
        collector = Collector(Config(Path("data")), Mock(), Mock(), llm)
        rows, warnings = collector.assess_relevance({"name": "金属导电理论", "query": "金属导电理论"}, records, queries, lambda _: None)
        self.assertFalse(warnings)
        self.assertEqual(relevance(rows[0], queries), 1)
        self.assertEqual(relevance(rows[1], queries), 0)
        self.assertEqual(relevance(rows[0], ["unrelated topic"]), 0)
        audit, _ = screen(rows, queries, {"level": "low"}, 10, year=2026)
        self.assertEqual(sum(a["selected"] for a in audit), 1)

    def test_incomplete_semantic_response_falls_back_without_partial_changes(self):
        llm = Mock(ready=True)
        llm.complete.return_value = {"items": []}
        collector = Collector(Config(Path("data")), Mock(), Mock(), llm)
        rows, warnings = collector.assess_relevance({"name": "topic", "query": "topic"}, [{"title": "topic"}], ["topic"], lambda _: None)
        self.assertEqual(len(warnings), 1)
        self.assertNotIn("semantic_relevance", rows[0])

    def test_excel_exports_numeric_if_rank_and_verified_source(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(folder)
            topic = store.create_topic("machine learning potentials")["id"]
            store.save_source(topic, {"title": "Machine learning potentials", "kind": "article", "venue": "The Journal of Chemical Physics", "year": "2025"})
            Collector(Config(Path(folder)), Mock(), store).rerank(topic)
            xlsx, _ = export_topic(store, topic)
            wb = load_workbook(io.BytesIO(xlsx))
            ws = wb["资料目录"]
            columns = {c.value: c.column for c in ws[1]}
            self.assertEqual(ws.cell(2, columns["期刊IF"]).value, 3.7)
            self.assertEqual(ws.cell(2, columns["IF年份"]).value, 2025)
            self.assertTrue(ws.cell(2, columns["IF核实来源"]).hyperlink)
            self.assertEqual(ws.cell(2, columns["重要度排名"]).value, 1)
            wb.close()


if __name__ == "__main__":
    unittest.main()
