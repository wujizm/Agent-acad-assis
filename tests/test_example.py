import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from research_agent.store import Store
from scripts.create_example import create_example


class ExampleTests(unittest.TestCase):
    def test_public_example_builds_offline_and_exports_traceable_records(self):
        with tempfile.TemporaryDirectory() as tmp, \
                patch("research_agent.network.Network.get", side_effect=AssertionError("offline example")), \
                patch("research_agent.intelligence.LLM.complete", side_effect=AssertionError("no model calls")):
            root = Path(tmp) / "new-demo"
            result = create_example(root)
            self.assertEqual((result["sources"], result["sources_with_text"], result["reports_generated"]), (8, 3, 0))
            self.assertGreater(result["evidence_count"], 0)
            store = Store(root)
            sources = store.sources(result["topic_id"])
            self.assertTrue(all(not s["text"] for s in sources if s["kind"] != "other"))
            notes = [s for s in sources if s["text"]]
            self.assertTrue(all("原创教学" in s["text"] and s.get("original_file") for s in notes))
            answer = json.loads((root / "offline-qa.json").read_text(encoding="utf-8"))
            self.assertEqual(answer["mode"], "search")
            self.assertTrue(all(e["source_id"] in {s["id"] for s in notes} for e in answer["evidence"]))
            with zipfile.ZipFile(result["bundle"]) as bundle:
                wb = load_workbook(io.BytesIO(bundle.read("资料目录.xlsx")))
                sheet = wb["资料目录"]
                self.assertEqual(sheet.max_row, 9)
                for row in range(2, 10):
                    self.assertIn(sheet.cell(row, 17).hyperlink.target, bundle.namelist())
                self.assertEqual(sum(isinstance(sheet.cell(row, 28).value, (int, float)) for row in range(2, 10)), 3)
                wb.close()

    def test_example_refuses_existing_data_without_changing_it(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            sentinel = root / "my-notes.txt"
            sentinel.write_text("keep my research", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "空目录"):
                create_example(root)
            self.assertEqual(sentinel.read_text(encoding="utf-8"), "keep my research")
            self.assertFalse((root / "library.sqlite3").exists())


if __name__ == "__main__":
    unittest.main()
