import base64
import io
import json
import tempfile
import threading
import unittest
import zipfile
from http.server import ThreadingHTTPServer
from pathlib import Path
from unittest.mock import Mock, patch

import requests
from openpyxl import load_workbook

from research_agent.collect import Collector, crossref_record, openalex_record, extract_file, trusted_domains, PageText
from research_agent.config import Config
from research_agent.export import export_topic
from research_agent.intelligence import ResearchAgent, LLM, retrieve, validate_citations
from research_agent.network import public_url
from research_agent.server import Application, make_handler
from research_agent.store import Store, chunks


def sample(**overrides):
    return {"title": "Grain boundary strengthening in alloys", "title_zh": "合金的晶界强化",
            "kind": "article", "authors": "A. Researcher", "venue": "Example Journal", "year": "2025",
            "doi": "10.1234/example", "url": "https://example.org/paper", "evidence_level": "abstract",
            "text": "晶界强化与晶粒尺寸相关。这是测试用的人工编写文本，不是实际研究证据。 Grain boundary strengthening depends on grain size.",
            **overrides}


class FakeLLM:
    ready = True

    def __init__(self, label):
        self.label = label

    def complete(self, prompt, json_output=False):
        if json_output:
            return {"title_zh": "合金的晶界强化", "title_en": "Grain boundary strengthening in alloys",
                    "summary": "本资料讨论晶界与晶粒尺寸的关系。", "main_content": "测试整理文本。"}
        if "检索词" in prompt:
            return "晶界 强化 grain boundary strengthening"
        return f"证据片段讨论晶界与晶粒尺寸的关系。[{self.label}]"


class StoreTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = Store(self.tmp.name)
        self.topic = self.store.create_topic("合金", "alloy")
        self.topic_id = self.topic["id"]

    def tearDown(self):
        self.tmp.cleanup()

    def add(self, **kwargs):
        return self.store.save_source(self.topic_id, sample(**kwargs))[0]

    def test_doi_dedup_and_upgrade_preserves_text(self):
        first = self.add()
        updated, created = self.store.save_source(self.topic_id, sample(doi="https://doi.org/10.1234/EXAMPLE",
            text="更长的全文内容 " * 200, evidence_level="fulltext"))
        self.assertFalse(created)
        self.assertEqual(first["id"], updated["id"])
        again = self.add(text="", evidence_level="metadata", affiliations="A. Researcher：University X")
        self.assertEqual(again["evidence_level"], "fulltext")
        self.assertEqual(len(self.store.sources(self.topic_id)), 1)
        self.assertIn("University X", again["affiliations"])
        self.assertTrue(self.store.evidence(self.topic_id))

    def test_different_book_editions_stay_distinct(self):
        self.add(kind="book", doi="", isbn="9781234567890", url="https://example.org/book/1")
        self.add(kind="book", doi="", isbn="9781234567891", url="https://example.org/book/2")
        self.assertEqual(len(self.store.sources(self.topic_id)), 2)

    def test_tracking_url_dedup(self):
        first = self.add(doi="", url="https://example.org/paper?utm_source=a")
        second = self.add(doi="", url="https://example.org/paper")
        self.assertEqual(first["id"], second["id"])

    def test_missing_affiliation_is_not_fabricated(self):
        source = self.add()
        self.assertEqual(source["affiliations"], "")
        card = (self.store.topic_dir(self.topic_id) / "sources" / f"{source['id']}.md").read_text(encoding="utf-8")
        self.assertIn("作者单位（来源明确提供）：未提供", card)

    def test_retrieval_isolation_exclusion_and_source_text(self):
        source = self.add()
        another = self.store.create_topic("另一主题")["id"]
        self.assertEqual(retrieve(self.store, another, "晶界强化"), [])
        self.assertTrue(retrieve(self.store, self.topic_id, "晶界强化"))
        self.assertTrue(retrieve(self.store, self.topic_id, "grain boundary"))
        source["main_content"] = "不应进入原始证据的术语supercalifragilistic"
        self.store.save_source(self.topic_id, source, source["id"])
        self.assertEqual(retrieve(self.store, self.topic_id, "supercalifragilistic"), [])
        source["excluded"] = True
        self.store.save_source(self.topic_id, source, source["id"])
        self.assertEqual(retrieve(self.store, self.topic_id, "晶界强化"), [])

    def test_report_citations_and_staleness(self):
        source = self.add()
        agent = ResearchAgent(self.store, FakeLLM(source["id"] + ":C1"))
        agent.report(self.topic_id, "review", lambda _: None)
        self.assertFalse(self.store.reports(self.topic_id)[0]["stale"])
        self.assertIn(source["url"], self.store.reports(self.topic_id)[0]["content"])
        source["notes"] = "修正了一条资料"
        self.store.save_source(self.topic_id, source, source["id"])
        self.assertTrue(self.store.reports(self.topic_id)[0]["stale"])

    def test_unknown_report_citation_not_persisted(self):
        self.add()
        with self.assertRaises(ValueError):
            ResearchAgent(self.store, FakeLLM("Saaaaaaaaaaaa:C999")).report(self.topic_id, "review", lambda _: None)
        self.assertEqual(self.store.reports(self.topic_id), [])

    def test_report_retrieves_relevant_late_sections_of_long_sources(self):
        source, _ = self.store.save_source(self.topic_id, sample(text=("Preface unrelated text. " * 400) + "\nAlloy grain boundary strengthening: late section evidence.", evidence_level="fulltext"))
        llm = Mock(ready=True)
        llm.complete.return_value = f"测试说明。[{source['id']}:C1]"
        ResearchAgent(self.store, llm).report(self.topic_id, "review", lambda _: None)
        self.assertIn("late section evidence", llm.complete.call_args.args[0])

    def test_reference_lists_do_not_replace_body_evidence(self):
        source = self.add(evidence_level="fulltext", text_version="author preprint v3")
        bibliography = "\n".join(f"{n}. Author, A. Alloy transport reference-only title. Journal 12, 34 (2020)." for n in range(11, 13))
        documents = [{"source_id": source["id"], "ordinal": i, "text": text} for i, text in enumerate([
            "Alloy introduction.", bibliography,
            "Alloy transport body evidence.\n1. Prepare specimen.\n2. Measure resistance.\n3. Compare results.",
            "Alloy transport substantive body. " * 30 + "\n" + bibliography.replace("reference-only title", "tail entry")], 1)]
        with patch.object(self.store, "evidence", return_value=documents):
            evidence = retrieve(self.store, self.topic_id, "alloy transport")
            self.assertEqual({d["ordinal"] for d in evidence}, {1, 3, 4})
            self.assertTrue(all(d["text_version"] == "author preprint v3" for d in evidence))
            llm = Mock(ready=True)
            llm.complete.return_value = f"测试说明。[{source['id']}:C3]"
            ResearchAgent(self.store, llm).report(self.topic_id, "review", lambda _: None)
            prompt = llm.complete.call_args.args[0]
            self.assertNotIn("reference-only title", prompt)
            self.assertIn("body evidence", prompt)

    def test_metadata_only_cannot_create_report(self):
        self.add(text="", evidence_level="metadata")
        with self.assertRaisesRegex(ValueError, "只有书目信息"):
            ResearchAgent(self.store, FakeLLM("unused")).report(self.topic_id, "review", lambda _: None)

    def test_offline_question_returns_evidence_not_ai_answer(self):
        self.add()
        result = ResearchAgent(self.store, LLM(Config(Path(self.tmp.name)))).ask(self.topic_id, "晶界强化")
        self.assertEqual(result["mode"], "search")
        self.assertTrue(result["evidence"])

    def test_ai_enrichment_and_answer(self):
        source = self.add()
        agent = ResearchAgent(self.store, FakeLLM(source["id"] + ":C1"))
        self.assertEqual(agent.enrich(self.topic_id, lambda _: None)["updated"], 1)
        self.assertEqual(agent.enrich(self.topic_id, lambda _: None)["updated"], 0)
        result = agent.ask(self.topic_id, "晶界强化")
        self.assertEqual(result["mode"], "ai")
        self.assertIn(source["id"], result["answer"])

    def test_editing_original_marks_ai_summary_stale(self):
        source = self.add()
        agent = ResearchAgent(self.store, FakeLLM(source["id"] + ":C1"))
        agent.enrich(self.topic_id, lambda _: None)
        source = self.store.source(source["id"])
        self.assertFalse(source["ai_stale"])
        source["text"] += "新增原文。"
        self.store.save_source(self.topic_id, source, source["id"])
        self.assertTrue(self.store.source(source["id"])["ai_stale"])

    def test_all_collection_failures_are_not_empty_success(self):
        network = Mock()
        network.json.side_effect = ValueError("测试连接失败")
        collector = Collector(Config(Path(self.tmp.name)), network, self.store)
        with self.assertRaisesRegex(ValueError, "所有搜集来源"):
            collector.run(self.topic_id, 1, lambda _: None)
        self.assertTrue((self.store.topic_dir(self.topic_id) / "last_search.json").is_file())

    def test_partial_collection_preserves_success_and_reports_failure(self):
        network = Mock()
        network.json.side_effect = [{"message": {"items": [{"title": ["Alloy study"], "DOI": "10.1/test"}]}},
                                    ValueError("近期采样失败"), ValueError("书籍源失败")]
        result = Collector(Config(Path(self.tmp.name)), network, self.store).run(self.topic_id, 1, lambda _: None)
        self.assertEqual(result["counts"]["Crossref"]["added"], 1)
        self.assertIn("error", result["counts"]["Open Library"])
        self.assertEqual(len(self.store.sources(self.topic_id)), 1)

    def test_excel_zip_links_and_formula_injection(self):
        source = self.add(title='=HYPERLINK("https://bad.example","x")')
        xlsx, archive = export_topic(self.store, self.topic_id)
        wb = load_workbook(io.BytesIO(xlsx))
        sheet = wb["资料目录"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["B2"].data_type, "s")
        self.assertEqual(sheet["B2"].value, source["title"])
        self.assertEqual(sheet.freeze_panes, "C2")
        with zipfile.ZipFile(io.BytesIO(archive)) as z:
            self.assertIn(sheet["Q2"].hyperlink.target, z.namelist())
            self.assertIn("资料目录.xlsx", z.namelist())
            self.assertIn(f"records/{source['id']}.json", z.namelist())
        wb.close()


class ParsingTests(unittest.TestCase):
    def test_crossref_author_affiliations(self):
        r = crossref_record({"title": ["Test"], "DOI": "10.1/ABC", "type": "journal-article", "publisher": "Publisher",
            "abstract": "<jats:p>Real abstract.</jats:p>", "author": [{"given": "Alice", "family": "Lee",
            "affiliation": [{"name": "Institution A"}]}, {"given": "Bob", "family": "Chan"}],
            "published": {"date-parts": [[2024]]}})
        self.assertEqual(r["affiliations"], "Alice Lee：Institution A")
        self.assertNotIn("Publisher", r["affiliations"])
        self.assertEqual(r["evidence_level"], "abstract")
        self.assertEqual(r["year"], "2024")

    def test_openalex_abstract_and_units(self):
        r = openalex_record({"id": "https://openalex.org/W1", "title": "Test", "abstract_inverted_index": {"World": [1], "Hello": [0]},
            "authorships": [{"author": {"display_name": "Alice"}, "institutions": [{"display_name": "University"}]}]})
        self.assertEqual(r["text"], "Hello World")
        self.assertEqual(r["affiliations"], "Alice：University")

    def test_html_excludes_script_and_gets_metadata(self):
        p = PageText()
        p.feed('<title>My paper</title><meta name="citation_author" content="Alice"><script>secret</script><p>Visible text</p>')
        self.assertNotIn("secret", p.text())
        self.assertIn("Visible text", p.text())
        self.assertEqual(p.meta["citation_author"], ["Alice"])

    def test_pdf_text_with_pages(self):
        import pymupdf
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Test fixture only. Grain boundary strengthening depends on the grain size in this example.")
        text, _ = extract_file(doc.tobytes(), "test.pdf")
        doc.close()
        self.assertIn("[第 1 页]", text)
        self.assertIn("Grain boundary", text)

    def test_blank_pdf_requires_ocr(self):
        import pymupdf
        doc = pymupdf.open()
        doc.new_page()
        with self.assertRaisesRegex(ValueError, "OCR"):
            extract_file(doc.tobytes(), "scan.pdf")
        doc.close()

    def test_chunks_cover_tail(self):
        text = "abc " * 1000 + "最后的证据"
        result = list(chunks(text))
        self.assertTrue(result[-1].endswith("最后的证据"))
        self.assertGreater(len(result), 1)

    def test_citation_validation(self):
        validate_citations("事实[S0123456789ab:C1]", ["S0123456789ab:C1"])
        for text in ("没有引用", "结论[S0123456789ab:C2]", "网址https://fake.test[S0123456789ab:C1]"):
            with self.assertRaises(ValueError):
                validate_citations(text, ["S0123456789ab:C1"])

    def test_private_source_urls_blocked(self):
        with self.assertRaises(ValueError):
            public_url("file:///etc/passwd")
        with patch("socket.getaddrinfo", return_value=[(2, 1, 6, "", ("127.0.0.1", 80))]):
            with self.assertRaises(ValueError):
                public_url("http://example.test")
        self.assertEqual(trusted_domains("nist.gov, https://www.energy.gov/"), ["nist.gov", "www.energy.gov"])


class HttpTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.app = Application(Config(Path(self.tmp.name)))
        self.server = ThreadingHTTPServer(("127.0.0.1", 0), make_handler(self.app))
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.base = f"http://127.0.0.1:{self.server.server_port}"
        self.headers = {"X-CSRF-Token": self.app.csrf}

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join()
        self.tmp.cleanup()

    def test_ui_and_full_local_http_flow(self):
        self.assertEqual(requests.get(self.base, timeout=5).status_code, 200)
        license_response = requests.get(self.base + "/license", timeout=5)
        self.assertEqual(license_response.status_code, 200)
        self.assertIn("GNU AFFERO GENERAL PUBLIC LICENSE", license_response.text)
        self.assertNotIn("llm_api_key", requests.get(self.base + "/api/status", timeout=5).json())
        topic = requests.post(self.base + "/api/topics", json={"name": "HTTP test"}, headers=self.headers, timeout=5).json()
        source = requests.post(self.base + f"/api/topics/{topic['id']}/sources", json=sample(), headers=self.headers, timeout=5).json()
        self.assertIn("id", source)
        bundle = requests.get(self.base + f"/api/topics/{topic['id']}/export?format=zip", timeout=5)
        self.assertTrue(zipfile.is_zipfile(io.BytesIO(bundle.content)))
        card = requests.get(self.base + f"/files/{topic['id']}/sources/{source['id']}.md", timeout=5)
        self.assertEqual(card.status_code, 200)

    def test_cross_origin_and_missing_csrf_rejected(self):
        self.assertEqual(requests.post(self.base + "/api/topics", json={"name": "bad"}, timeout=5).status_code, 403)
        self.assertEqual(requests.post(self.base + "/api/topics", json={"name": "bad"},
            headers={**self.headers, "Origin": "https://evil.example"}, timeout=5).status_code, 403)

    def test_uploaded_raw_file_is_retained(self):
        topic = self.app.store.create_topic("Upload")
        data = b"Test material notes with enough text to search and preserve as the original file."
        result = self.app.upload(topic["id"], {"filename": "notes.txt", "content": base64.b64encode(data).decode()}, lambda _: None)
        source = self.app.store.source(result["source_id"])
        self.assertEqual((self.app.store.topic_dir(topic["id"]) / source["original_file"]).read_bytes(), data)
        again = self.app.upload(topic["id"], {"filename": "notes.txt", "content": base64.b64encode(data).decode()}, lambda _: None)
        self.assertEqual(again["source_id"], source["id"])
        self.assertFalse(again["created"])

    def test_model_settings_and_connection_job_require_csrf(self):
        body = {"llm_backend": "codex", "codex_model": "test-model"}
        for path in ("/api/model-settings", "/api/model-test"):
            self.assertEqual(requests.post(self.base + path, json=body, timeout=5).status_code, 403)
        saved = requests.post(self.base + "/api/model-settings", json=body, headers=self.headers, timeout=5)
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.json()["llm_backend"], "codex")
        with patch.object(self.app.agent.llm, "check_connection", return_value={"answer": "连接成功"}) as check:
            response = requests.post(self.base + "/api/model-test", json={}, headers=self.headers, timeout=5)
            self.assertEqual(response.status_code, 200)
            for _ in range(50):
                job = self.app.jobs.get(response.json()["job_id"])
                if job["state"] != "running":
                    break
                threading.Event().wait(.01)
            self.assertEqual(job["state"], "done")
            self.assertEqual(job["result"]["answer"], "连接成功")
            check.assert_called_once()


class ModelAdapterTests(unittest.TestCase):
    def test_compatible_provider_response_and_request(self):
        config = Config(Path("unused"), llm_base_url="https://provider.example/v1", llm_model="test-model", llm_api_key="test-only")
        response = Mock(status_code=200)
        response.json.return_value = {"choices": [{"finish_reason": "stop", "message": {"content": '{"summary":"测试输出"}'}}]}
        with patch("research_agent.intelligence.requests.post", return_value=response) as post:
            result = LLM(config).complete("只使用测试夹具", json_output=True)
        self.assertEqual(result["summary"], "测试输出")
        self.assertEqual(post.call_args.args[0], "https://provider.example/v1/chat/completions")
        self.assertEqual(post.call_args.kwargs["json"]["model"], "test-model")

    def test_truncated_response_is_rejected(self):
        config = Config(Path("unused"), llm_base_url="http://localhost:11434/v1", llm_model="test-model")
        response = Mock(status_code=200)
        response.json.return_value = {"choices": [{"finish_reason": "length", "message": {"content": "不完整的内容"}}]}
        with patch("research_agent.intelligence.requests.post", return_value=response):
            with self.assertRaisesRegex(ValueError, "截断"):
                LLM(config).complete("测试")


if __name__ == "__main__":
    unittest.main()
