import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from research_agent.collect import Collector, crossref_record, openalex_record
from research_agent.config import Config
from research_agent.export import export_topic
from research_agent.screening import assess_heat, exclusion, relevance, screen
from research_agent.store import Store


def paper(index, count=0, year=2026, **changes):
    return {"title": f"AI for science: study {index}", "kind": "article", "year": str(year),
            "doi": f"10.1234/test{index}", "url": f"https://doi.org/10.1234/test{index}",
            "text": "Artificial intelligence for scientific discovery. Test fixture only.", "evidence_level": "abstract",
            "citation_metrics": {"count": count, "provider": "Crossref", "retrieved_at": "2026-10-17"}, **changes}


class ScreeningTests(unittest.TestCase):
    def test_heat_uses_relevant_recent_unique_research_not_raw_hits(self):
        records = [paper(i) for i in range(40)]
        records += [paper(i) for i in range(40)]  # Duplicate provider/DOI.
        records += [paper(80, year=2000), paper(81, title="Road slope protection", text=""),
                    paper(82, record_type="peer-review"), paper(83, kind="book")]
        heat = assess_heat(records, "AI for Science", ["Crossref"], 2026)
        self.assertEqual(heat["recent_count"], 40)
        self.assertEqual(heat["level"], "high")
        self.assertEqual(assess_heat([paper(i) for i in range(15)], "AI for Science", ["Crossref"], 2026)["level"], "medium")
        self.assertEqual(assess_heat([paper(1)], "AI for Science", ["Crossref"], 2026)["level"], "low")
        self.assertEqual(assess_heat([], "AI for Science", [], 2026)["level"], "unknown")

    def test_relatedness_uses_title_and_abstract_not_authors_or_journal(self):
        self.assertEqual(relevance(paper(1, title="Road slope protection", text="", authors="AI Science", venue="AI for Science"), "AI for Science"), 0)
        self.assertGreater(relevance(paper(1, title="Artificial intelligence for scientific discovery"), "AI for Science"), 0.65)
        self.assertGreater(relevance({"title": "High entropy alloys", "text": ""}, "high entropy alloy"), 0.65)
        self.assertGreater(relevance({"title": "高熵合金的强化机制", "text": ""}, "高熵合金强化"), 0.65)

    def test_retracted_notices_peer_reviews_and_future_years_are_deferred(self):
        for changes in ({"title": "(WITHDRAWN) AI for Science"}, {"title": "Correction to: AI for Science"},
                        {"is_retracted": True}, {"record_type": "peer-review"}, {"year": "2028"}):
            self.assertTrue(exclusion(paper(1, **changes), "AI for Science", 2026))
        self.assertEqual(exclusion(paper(1, title="Retraction detection with AI for science"), "AI for Science", 2026), "")

    def test_hot_topic_prefers_attention_and_preserves_zero_citation_frontier(self):
        records = [paper(i, count=300 + i * 10, year=2020) for i in range(20)]
        records += [paper(21, count=0, year=2026), paper(22, count=None, year=2025), paper(23, count=0, year=2010)]
        audit, policy = screen(records, "AI for Science", {"level": "high"}, 10, year=2026)
        chosen = [d for d in audit if d["selected"]]
        self.assertGreaterEqual(policy["annual_citation_threshold"], 2)
        self.assertTrue(any(d["record"]["doi"] == "10.1234/test21" and d["path"] == "近期研究" for d in chosen))
        self.assertTrue(any(d["record"]["doi"] == "10.1234/test22" and d["path"] == "近期研究" for d in chosen))
        self.assertFalse(any(d["record"]["doi"] == "10.1234/test23" for d in chosen))
        self.assertLessEqual(len(chosen), 10)

    def test_low_heat_and_override_are_less_restrictive(self):
        records = [paper(i, count=0, year=2010) for i in range(6)]
        for heat, mode in (("low", "auto"), ("unknown", "auto"), ("high", "broad")):
            audit, policy = screen(records, "AI for Science", {"level": heat}, 10, mode, 2026)
            self.assertEqual(sum(d["selected"] for d in audit), 6)
            self.assertIsNone(policy["annual_citation_threshold"])
        audit, _ = screen(records, "AI for Science", {"level": "low"}, 10, "strict", 2026)
        self.assertFalse(any(d["selected"] for d in audit))

    def test_books_do_not_require_citations_and_missing_is_not_zero(self):
        audit, policy = screen([paper(1, count=None, kind="book", year=2000)], "AI for Science", {"level": "high"}, 10, year=2026)
        self.assertTrue(audit[0]["selected"])
        self.assertIsNone(audit[0]["annual_citations"])
        self.assertIsNone(policy["annual_citation_threshold"])
        self.assertEqual(crossref_record({"title": ["Test"], "is-referenced-by-count": 0})["citation_metrics"]["count"], 0)
        self.assertIsNone(crossref_record({"title": ["Test"]})["citation_metrics"]["count"])
        self.assertEqual(openalex_record({"id": "https://openalex.org/W1", "title": "Test", "cited_by_count": 9})["citation_metrics"]["provider"], "OpenAlex")

    def test_invalid_mode_is_rejected(self):
        with self.assertRaises(ValueError):
            screen([], "test", {"level": "low"}, 10, "invalid", 2026)

    def test_possible_versions_keep_separate_audit_but_use_one_slot(self):
        records = [paper(1, kind="chapter"), paper(2, kind="chapter", title="AI for science: study 1"), paper(3, kind="book")]
        audit, _ = screen(records, "AI for Science", {"level": "high"}, 10, year=2026)
        self.assertEqual(len(audit), 3)
        self.assertEqual(sum(d["selected"] for d in audit), 2)
        self.assertIn("可能版本", next(d["reason"] for d in audit if not d["selected"]))

    def test_frontier_balances_attention_after_relevance_gate(self):
        records = [paper(i, count=700 + i * 10, year=2020) for i in range(8)]
        records += [paper(i, count=10, title=f"Scientific discovery in autonomous laboratories {i}") for i in range(10, 13)]
        records += [paper(20, count=0, title="Artificial intelligence for scientific discovery")]
        audit, _ = screen(records, "artificial intelligence scientific discovery", {"level": "high"}, 10, year=2026)
        self.assertEqual(sum(d["path"] == "近期研究" for d in audit), 3)
        self.assertFalse(next(d["selected"] for d in audit if d["record"]["doi"] == "10.1234/test20"))


class ScreeningIntegrationTests(unittest.TestCase):
    def test_collection_audit_manual_decisions_and_export_roundtrip(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(folder)
            topic = store.create_topic("AI for Science")["id"]
            incoming = {"title": ["AI for Science"], "type": "journal-article", "DOI": "10.1/test",
                        "published": {"date-parts": [[2026]]}, "is-referenced-by-count": 0}
            raw = {"message": {"total-results": 9000000, "items": [incoming,
                   {**incoming, "title": ["Review for AI for Science"], "DOI": "10.1/peer", "type": "peer-review"}]}}
            net = Mock()
            net.json.side_effect = [raw, raw, {"docs": []}] * 2
            collector = Collector(Config(Path(folder)), net, store)
            with patch("research_agent.collect.date") as today:
                from datetime import date
                today.today.return_value = date(2026, 10, 17)
                first = collector.run(topic, 10, lambda _: None)
                self.assertEqual(first["heat"]["recent_count"], 1)
                source = store.sources(topic)[0]
                source.update(excluded=True, verified=True)
                store.save_source(topic, source, source["id"])
                second = collector.run(topic, 10, lambda _: None)
            self.assertEqual(second["counts"]["Crossref"]["merged"], 1)
            self.assertTrue(store.sources(topic)[0]["excluded"])
            self.assertTrue(store.sources(topic)[0]["verified"])
            self.assertNotEqual(first["audit_file"], second["audit_file"])
            audit = json.loads((store.topic_dir(topic) / second["audit_file"]).read_text(encoding="utf-8"))
            self.assertEqual(len(audit["decisions"]), 2)
            self.assertEqual(sum(d["selected"] for d in audit["decisions"]), 1)
            self.assertEqual(store.last_search(topic)["mode"], "auto")
            xlsx, archive = export_topic(store, topic)
            wb = load_workbook(io.BytesIO(xlsx))
            ws = wb["资料目录"]
            self.assertEqual(ws["T2"].value, "0")
            self.assertEqual(ws["U2"].value, "Crossref")
            self.assertTrue(ws["Y2"].value)
            self.assertEqual(ws.auto_filter.ref, "A1:AG2")
            self.assertEqual(ws["P2"].value, "否")
            import zipfile
            with zipfile.ZipFile(io.BytesIO(archive)) as z:
                self.assertIn(first["audit_file"], z.namelist())
                card = z.read(ws["Q2"].hyperlink.target).decode("utf-8")
                self.assertIn("引用次数：0", card)
                self.assertIn("入选依据", card)
            wb.close()

    def test_metric_source_and_zero_update_together_on_merge(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(folder)
            topic = store.create_topic("test")["id"]
            store.save_source(topic, paper(1, count=100))
            incoming = paper(1, count=0)
            incoming["citation_metrics"]["provider"] = "OpenAlex"
            source, created = store.save_source(topic, incoming)
            self.assertFalse(created)
            self.assertEqual(source["citation_metrics"]["count"], 0)
            self.assertEqual(source["citation_metrics"]["provider"], "OpenAlex")


if __name__ == "__main__":
    unittest.main()
