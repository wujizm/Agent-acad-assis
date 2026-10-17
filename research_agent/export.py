import io
import json
import re
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from .store import LEVELS, atomic_text


COLUMNS = [("资料ID", "id", 20), ("资料名称（原题名）", "title", 55), ("题名（中文）", "title_zh", 55),
           ("题名（英文）", "title_en", 55), ("来源（期刊/出版社/机构）", "venue", 32), ("作者", "authors", 35),
           ("作者主要单位（原始关联）", "affiliations", 45), ("一句话简介", "summary", 65), ("资料类型", "kind", 15),
           ("年份", "year", 12), ("DOI", "doi", 32), ("ISBN", "isbn", 20), ("网址", "url", 40),
           ("内容覆盖范围", "evidence_level", 22), ("人工核验", "verified", 14), ("纳入综述", "included", 14),
           ("独立资料文件", "card", 28), ("获取日期UTC", "created_at", 26), ("备注", "notes", 60),
           ("引用次数（非影响因子）", "citation_count", 22), ("引用数据来源", "citation_provider", 20),
           ("引用数据时间UTC", "citation_date", 26), ("候选集热度", "heat", 16),
           ("入选通道", "screening_path", 18), ("入选依据", "screening_reason", 65),
           ("重要度排名", "importance_rank", 16), ("阅读重要度分", "importance_score", 18),
           ("期刊IF", "impact_factor", 14), ("IF年份", "jif_year", 12),
           ("IF核实状态", "jif_status", 16), ("IF核实来源", "jif_source", 48),
           ("IF核实日期", "jif_date", 20), ("重要度依据", "importance_reason", 75)]
COMPACT_COLUMNS = {"importance_rank", "importance_score", "impact_factor", "jif_year", "jif_status", "jif_date"}


def cell_text(value):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value if value is not None else ""))[:32767]


def workbook_bytes(store, topic_id):
    sources = sorted(store.sources(topic_id), key=lambda s: (bool(s["excluded"]), s.get("importance", {}).get("rank") or 999999))
    wb = Workbook()
    ws = wb.active
    ws.title = "资料目录"
    ws.append([c[0] for c in COLUMNS])
    for i, source in enumerate(sources, 2):
        metric, screening = source.get("citation_metrics", {}), source.get("screening", {})
        journal, priority = source.get("journal_metrics", {}), source.get("importance", {})
        values = {**source, "evidence_level": LEVELS[source["evidence_level"]],
                  "verified": "已核验" if source["verified"] else "待核验", "included": "否" if source["excluded"] else "是",
                  "card": f"sources/{source['id']}.md", "citation_count": metric.get("count"),
                  "citation_provider": metric.get("provider"), "citation_date": metric.get("retrieved_at"),
                  "heat": screening.get("heat"), "screening_path": screening.get("path"), "screening_reason": screening.get("reason"),
                  "importance_rank": priority.get("rank"), "importance_score": priority.get("score"),
                  "impact_factor": journal.get("value"), "jif_year": journal.get("year"),
                  "jif_status": journal.get("status", "未核实"), "jif_source": journal.get("source_url"),
                  "jif_date": journal.get("retrieved_at"), "importance_reason": priority.get("reason")}
        for j, (_, key, _) in enumerate(COLUMNS, 1):
            cell = ws.cell(i, j, cell_text(values.get(key)) or "未提供")
            cell.data_type = "s"  # Imported titles beginning with '=' must remain text.
            if key in ("importance_rank", "importance_score", "impact_factor", "jif_year") and isinstance(values.get(key), (int, float)):
                cell.value = values[key]
                cell.number_format = "0.0" if key in ("importance_score", "impact_factor") else "0"
            cell.alignment = Alignment(horizontal="center" if key in COMPACT_COLUMNS else None, vertical="top", wrap_text=True)
            if key in ("url", "card", "jif_source") and values.get(key):
                cell.hyperlink = values[key]
                cell.font = Font(name="Microsoft YaHei", color="146C94", underline="single")
            else:
                cell.font = Font(name="Microsoft YaHei", size=10)
        ws.row_dimensions[i].height = 62
    for j, (_, key, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(j)].width = width
        cell = ws.cell(1, j)
        cell.fill = PatternFill("solid", fgColor="003366")
        cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center" if key in COMPACT_COLUMNS else None, wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "C2"
    ws.sheet_view.showGridLines = False
    last_column = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A1:{last_column}{max(1, ws.max_row)}"
    if sources:
        table = Table(displayName="Sources", ref=f"A1:{last_column}{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    notes = wb.create_sheet("阅读说明")
    for row in [
        ("主题", store.topic(topic_id)["name"]),
        ("字段约定", "题名是资料题名；来源单列为期刊/出版社/机构。译名与简介可能包含 AI 整理，应对照原文。"),
        ("缺失信息", "未提供表示来源中缺失，作者单位不得依据常识补造。"),
        ("独立文件", "相对链接指向 sources/ 下的 Markdown 文件。请下载整个主题 ZIP 并解压后使用。"),
        ("数据编辑", "在应用内编辑资料后重新导出。Excel/Markdown 的外部编辑不会同步回数据库。"),
        ("内容范围", "仅书目、摘要、节选与全文必须区分；完整元数据及提取文本在 records/ 与 sources/ 中。"),
        ("自动筛选", "引用数为指定来源/日期的记录，不是期刊影响因子；年均引用仅作本来源内排序。候选集热度不是全领域统计，未入选候选及理由见 searches/。"),
        ("IF与重要度", "IF附指标年份及官方核实来源；缺失不记为0，期刊IF不是单篇论文质量。重要度：相关性60/同来源论文关注度20/IF10/原文覆盖5/时效5，缺失项重分权重；仅作本主题阅读顺序。"),
        ("报告", "生成结果是 AI 草稿；修改资料后需要重新生成报告。")]:
        notes.append(row)
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 100
    for row in notes:
        for cell in row:
            cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        notes.row_dimensions[row[0].row].height = 36
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def export_topic(store, topic_id):
    with store.lock:
        root = store.topic_dir(topic_id)
        sources = store.sources(topic_id)
        for source in sources:
            store.write_card(source)
        xlsx = workbook_bytes(store, topic_id)
        (root / "资料目录.xlsx").write_bytes(xlsx)
        atomic_text(root / "topic.json", json.dumps(store.topic(topic_id), ensure_ascii=False, indent=2))
        for report in store.reports(topic_id):
            prefix = "> 注意：资料已更新，此报告需要重新生成。\n\n" if report["stale"] else ""
            atomic_text(root / "reports" / f"{report['kind']}.md", prefix + report["content"])
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            for path in sorted(root.rglob("*")):
                if path.is_file() and path.suffix not in (".zip", ".tmp"):
                    archive.write(path, path.relative_to(root).as_posix())
        return xlsx, output.getvalue()
